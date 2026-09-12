"""Word (.docx) and PDF report generation.

Spec addition (V1-0b). No Streamlit imports.

Both libraries are optional: if python-docx or reportlab is not installed the
relevant function raises ReportUnavailable and the UI hides that button rather
than crashing.

Note on the PDF: reportlab's built-in fonts are Latin-1 only, so all text going
into the PDF is passed through _ascii() — chemical formulae are written Ta2O5
rather than with subscripts, and dashes are plain hyphens.
"""
from __future__ import annotations

import io
import re
from datetime import datetime

import pandas as pd

BRAND_GREEN = "1A5C38"
BRAND_GOLD = "C8960A"

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
    )
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False


class ReportUnavailable(Exception):
    """The library needed for this format is not installed."""


def _ascii(text) -> str:
    """Latin-1-safe text for reportlab's built-in fonts."""
    s = str(text)
    for bad, good in (("\u2014", "-"), ("\u2013", "-"), ("\u2018", "'"),
                      ("\u2019", "'"), ("\u201c", '"'), ("\u201d", '"'),
                      ("\u2192", "->"), ("\u00d7", "x"), ("\u2264", "<="),
                      ("\u2265", ">="), ("\u2082", "2"), ("\u2085", "5")):
        s = s.replace(bad, good)
    return re.sub(r"[^\x00-\xff]", "?", s)


def _fmt_cell(v, col: str = "") -> str:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return "-"
    # Years must never be comma-grouped: 2024, not 2,024.
    is_year_col = "year" in str(col).lower() or str(col).lower() in ("period", "yr")
    if isinstance(v, (int,)) or (isinstance(v, float) and float(v).is_integer()):
        iv = int(v)
        if is_year_col or 1900 <= iv <= 2200:
            return str(iv)
    if isinstance(v, float):
        return f"{v:,.2f}" if abs(v) < 1000 else f"{v:,.0f}"
    if isinstance(v, (int,)):
        return f"{v:,}"
    return str(v)


# ------------------------------------------------------------------- DOCX ----

def build_docx(context: dict, tables: dict[str, pd.DataFrame]) -> bytes:
    """Word report. `context` carries the headline figures and caveats."""
    if not DOCX_AVAILABLE:
        raise ReportUnavailable(
            "python-docx is not installed. Add `python-docx>=1.1` to "
            "requirements.txt and reinstall."
        )

    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(0.8)

    title = doc.add_paragraph()
    run = title.add_run("Rwanda 3T Mineral Export Intelligence")
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(BRAND_GREEN)
    run.font.name = "Cambria"

    sub = doc.add_paragraph()
    srun = sub.add_run(context.get("subtitle", "Analysis report"))
    srun.font.size = Pt(11)
    srun.font.color.rgb = RGBColor.from_string(BRAND_GOLD)
    srun.font.name = "Cambria"

    meta = doc.add_paragraph()
    meta.add_run(
        f"Generated {datetime.now().strftime('%d-%b-%Y %H:%M')}  |  "
        f"Period {context.get('period', '-')}  |  "
        f"Filters: {context.get('filters', 'none')}"
    ).font.size = Pt(8)

    def _heading(text: str) -> None:
        """Heading in brand green - Word's built-in Heading 1 is blue."""
        h = doc.add_heading(text, level=1)
        for r in h.runs:
            r.font.color.rgb = RGBColor.from_string(BRAND_GREEN)
            r.font.name = "Cambria"

    # --- headline figures ---
    _heading("Headline figures")
    metrics = context.get("metrics", {})
    if metrics:
        t = doc.add_table(rows=1, cols=2)
        t.style = "Light Grid Accent 1"
        hdr = t.rows[0].cells
        hdr[0].text, hdr[1].text = "Measure", "Value"
        for k, v in metrics.items():
            cells = t.add_row().cells
            cells[0].text = str(k)
            cells[1].text = str(v)

    # --- tables ---
    for name, df in tables.items():
        if df is None or df.empty:
            continue
        _heading(str(name))
        shown = df.head(40).reset_index() if df.index.name else df.head(40)
        cols = list(shown.columns)[:9]

        t = doc.add_table(rows=1, cols=len(cols))
        t.style = "Light Grid Accent 1"
        for i, c in enumerate(cols):
            cell = t.rows[0].cells[i]
            cell.text = str(c)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.bold = True
                    r.font.size = Pt(8)
        for _, row in shown.iterrows():
            cells = t.add_row().cells
            for i, c in enumerate(cols):
                cells[i].text = _fmt_cell(row[c], c)
                for p in cells[i].paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(8)
        if len(df) > 40:
            note = doc.add_paragraph(
                f"Showing the first 40 of {len(df):,} rows. "
                "Use the Excel export for the complete table.")
            note.runs[0].font.size = Pt(8)
            note.runs[0].font.italic = True

    # --- caveats ---
    caveats = context.get("caveats", [])
    if caveats:
        doc.add_page_break()
        _heading("Caveats and basis")
        for c in caveats:
            p = doc.add_paragraph(str(c), style="List Bullet")
            p.runs[0].font.size = Pt(9)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# -------------------------------------------------------------------- PDF ----

def build_pdf(context: dict, tables: dict[str, pd.DataFrame]) -> bytes:
    if not PDF_AVAILABLE:
        raise ReportUnavailable(
            "reportlab is not installed. Add `reportlab>=4.0` to "
            "requirements.txt and reinstall."
        )

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4),
        leftMargin=1.2 * cm, rightMargin=1.2 * cm,
        topMargin=1.2 * cm, bottomMargin=1.2 * cm,
        title="Rwanda 3T Mineral Export Intelligence",
    )
    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontSize=18,
                        textColor=colors.HexColor(f"#{BRAND_GREEN}"),
                        spaceAfter=6)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontSize=12,
                        textColor=colors.HexColor(f"#{BRAND_GREEN}"),
                        spaceBefore=12, spaceAfter=4)
    small = ParagraphStyle("small", parent=ss["Normal"], fontSize=7.5,
                           textColor=colors.HexColor("#555555"))
    body = ParagraphStyle("body", parent=ss["Normal"], fontSize=8.5,
                          spaceAfter=3)
    # Paragraph objects carry their own colour and ignore a TableStyle
    # TEXTCOLOR, so header cells need their own white style.
    hdr = ParagraphStyle("hdr", parent=ss["Normal"], fontSize=7.5,
                         textColor=colors.white, fontName="Helvetica-Bold")

    story = [
        Paragraph(_ascii("Rwanda 3T Mineral Export Intelligence"), h1),
        Paragraph(_ascii(context.get("subtitle", "Analysis report")),
                  ParagraphStyle("sub", parent=ss["Normal"], fontSize=10,
                                 textColor=colors.HexColor(f"#{BRAND_GOLD}"))),
        Spacer(1, 4),
        Paragraph(_ascii(
            f"Generated {datetime.now().strftime('%d-%b-%Y %H:%M')}  |  "
            f"Period {context.get('period', '-')}  |  "
            f"Filters: {context.get('filters', 'none')}"), small),
        Spacer(1, 10),
    ]

    def _table(df: pd.DataFrame, max_rows: int = 28) -> Table:
        shown = df.head(max_rows).reset_index() if df.index.name else df.head(max_rows)
        cols = list(shown.columns)[:9]
        data = [[Paragraph(_ascii(c), hdr) for c in cols]]
        for _, row in shown.iterrows():
            data.append([Paragraph(_ascii(_fmt_cell(row[c], c)), small) for c in cols])
        t = Table(data, repeatRows=1, hAlign="LEFT")
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{BRAND_GREEN}")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CCCCCC")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1),
             [colors.white, colors.HexColor("#F2FAF5")]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ("TOPPADDING", (0, 0), (-1, -1), 2),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ]))
        return t

    metrics = context.get("metrics", {})
    if metrics:
        story.append(Paragraph("Headline figures", h2))
        story.append(_table(pd.DataFrame(
            {"Measure": list(metrics), "Value": [str(v) for v in metrics.values()]})))

    for name, df in tables.items():
        if df is None or df.empty:
            continue
        story.append(Paragraph(_ascii(str(name)), h2))
        story.append(_table(df))
        if len(df) > 28:
            story.append(Paragraph(_ascii(
                f"Showing the first 28 of {len(df):,} rows. "
                "Use the Excel export for the complete table."), small))

    caveats = context.get("caveats", [])
    if caveats:
        story.append(PageBreak())
        story.append(Paragraph("Caveats and basis", h2))
        for c in caveats:
            story.append(Paragraph(_ascii(f"- {c}"), body))

    doc.build(story)
    return buf.getvalue()


def standard_caveats() -> list[str]:
    """The caveats that must travel with any exported figure."""
    return [
        "All prices are US dollars per kilogram of CONTAINED metal "
        "(value divided by quantity x grade / 100), so grade differences do "
        "not distort any comparison.",
        "International benchmark reliability is not equal across the three "
        "minerals. Cassiterite uses LME tin cash settlement monthly averages "
        "and is exchange-traded and continuous. Coltan and wolframite "
        "concentrate do not trade on any exchange: the record is USGS annual "
        "averages plus a few spot points, and several 2026 months rest on a "
        "flat-held assumption that is very likely too low.",
        "Every benchmark figure carries a provenance tag (REAL, INTERPOLATED, "
        "ASSUMPTION or USER OVERRIDE). Check the tag before quoting a number.",
        "Annual benchmarks applied to monthly data understate within-year price "
        "swings.",
        "Loss figures are floored at zero: months where Rwanda's realised price "
        "exceeded the benchmark are treated as no loss rather than as a gain. "
        "This is the conservative choice for reporting.",
        "Peer comparison groups exporters shipping the SAME PRODUCT in the same "
        "period, not the same mineral.",
        "Trading below the peer average is NOT by itself evidence of "
        "wrongdoing. Contract timing, impurity penalties, lot size, payment "
        "terms and long-term offtake agreements priced off an earlier reference "
        "all produce legitimate discounts. This analysis identifies where to "
        "look, not what was found.",
        "Peer periods with fewer than three active exporters are flagged as "
        "thin: with one or two exporters the market is largely the exporter "
        "itself.",
        "Excluded records are governed by the exclusion rules in force at export "
        "time. Review them on the Reports page before relying on any total.",
    ]

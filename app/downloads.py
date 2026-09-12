"""Excel export helper — FinIntel R14.

Every table and chart section gets a download button. Workbooks are formatted
(frozen header, bold headers, auto-width) and carry an _INFO sheet recording
where the export came from.
"""
from __future__ import annotations

import io

import pandas as pd
import streamlit as st
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

TOOL_VERSION = "V1-0c"
HEADER_FILL = PatternFill("solid", start_color="1A5C38")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)


def stamp() -> str:
    """__V{yyyy-mm-dd}_{hhmm} — R14.3."""
    return pd.Timestamp.now().strftime("__V%Y-%m-%d_%H%M")


def build_workbook(sheets: dict[str, pd.DataFrame], source: str = "",
                   filters: str = "") -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xl:
        for name, df in sheets.items():
            safe = str(name)[:31]
            out = df.reset_index() if df.index.name or isinstance(
                df.index, pd.MultiIndex) else df
            out.to_excel(xl, sheet_name=safe, index=False)

        info = pd.DataFrame({
            "Field": ["Exported at", "Tool version", "Source", "Filters applied",
                      "Sheets"],
            "Value": [
                pd.Timestamp.now().strftime("%d-%b-%Y %H:%M"),
                TOOL_VERSION,
                source or "RMB 3T Export Intelligence",
                filters or "none",
                ", ".join(str(s)[:31] for s in sheets),
            ],
        })
        info.to_excel(xl, sheet_name="_INFO", index=False)

        for ws in xl.book.worksheets:
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.font = HEADER_FONT
                cell.fill = HEADER_FILL
                cell.alignment = Alignment(horizontal="center", vertical="center",
                                           wrap_text=True)
            for col in ws.columns:
                width = max(
                    (len(str(c.value)) for c in col[:300] if c.value is not None),
                    default=10,
                )
                ws.column_dimensions[get_column_letter(col[0].column)].width = \
                    min(max(width + 2, 10), 55)
    return buf.getvalue()


def download_button(label: str, sheets: dict[str, pd.DataFrame], basename: str,
                    source: str = "", filters: str = "", key: str | None = None,
                    width: str = "content") -> None:
    """R14.1 — a download button beside a table or chart section."""
    usable = {k: v for k, v in sheets.items()
              if isinstance(v, pd.DataFrame) and not v.empty}
    if not usable:
        return
    try:
        data = build_workbook(usable, source=source, filters=filters)
    except Exception as exc:  # noqa: BLE001 — never let an export break a page
        st.caption(f"Export unavailable: {exc}")
        return
    st.download_button(
        f"📥 DOWNLOAD — {label} (.xlsx)",
        data,
        file_name=f"{basename}{stamp()}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        key=key,
        width=width,
    )
